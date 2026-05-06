from pptx import Presentation
from pptx.util import Pt
from pptx.enum.text import PP_ALIGN

prs = Presentation(r"D:\spring-legacy\ticket-parent\tr(티레).pptx")

w = prs.slide_width
h = prs.slide_height
print(f"Slide size: {w} EMU x {h} EMU = {w/914400:.2f}\" x {h/914400:.2f}\"")
print(f"Total slides: {len(prs.slides)}")
print()

for i, slide in enumerate(prs.slides):
    print(f"=== Slide {i+1} ===")
    layout = slide.slide_layout
    print(f"  Layout: {layout.name}")
    bg = slide.background
    fill = bg.fill
    print(f"  BG fill type: {fill.type}")
    if fill.type is not None:
        try:
            print(f"  BG color: #{fill.fore_color.rgb}")
        except:
            pass
    print(f"  Shapes ({len(slide.shapes)}):")
    for j, shape in enumerate(slide.shapes):
        print(f"    [{j}] {shape.shape_type} name='{shape.name}' pos=({shape.left},{shape.top}) size=({shape.width},{shape.height})")
        if shape.has_text_frame:
            for para in shape.text_frame.paragraphs:
                for run in para.runs:
                    font = run.font
                    print(f"         text='{run.text[:60]}' font={font.name} size={font.size} bold={font.bold}")
                    if font.color.type is not None:
                        try:
                            print(f"         color=#{font.color.rgb}")
                        except:
                            pass
